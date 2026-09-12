"""DocxRenderer: renders a flattened doc-model tree into a themed Word
document, including headers/footers, page numbers, and inline screenshot
figures. See this package's __init__.py for the sibling list."""
import os
import logging

from docx import Document
from docx.shared import Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ..config import IMAGE_DOCX_WIDTH_INCHES

logger = logging.getLogger(__name__)


class DocxRenderer:
    def _add_page_number(self, run):
        fldChar1 = OxmlElement('w:fldChar')
        fldChar1.set(qn('w:fldCharType'), 'begin')
        instrText = OxmlElement('w:instrText')
        instrText.set(qn('xml:space'), 'preserve')
        instrText.text = "PAGE"
        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'separate')
        fldChar3 = OxmlElement('w:fldChar')
        fldChar3.set(qn('w:fldCharType'), 'end')
        run._r.append(fldChar1)
        run._r.append(instrText)
        run._r.append(fldChar2)
        run._r.append(fldChar3)

    def render(self, flat_model, out_path, context):
        doc = Document()
        is_dark = context.get('theme', '').lower() == 'dark'
        if is_dark:
            # Set background color
            background = OxmlElement('w:background')
            background.set(qn('w:color'), '121212')
            doc.element.insert(0, background)
            
            if not doc.settings.element.xpath('./w:displayBackgroundShape'):
                disp = OxmlElement('w:displayBackgroundShape')
                doc.settings.element.append(disp)
                
            # Set text colors for styles
            for style_name in ['Normal', 'List Paragraph', 'List Bullet']:
                try:
                    style = doc.styles[style_name]
                    if hasattr(style, 'font'):
                        style.font.color.rgb = RGBColor(0xE0, 0xE0, 0xE0)
                except KeyError:
                    pass
            for i in range(1, 4):
                try:
                    style = doc.styles[f'Heading {i}']
                    if hasattr(style, 'font'):
                        style.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                except KeyError:
                    pass
            try:
                doc.styles['Title'].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            except KeyError:
                pass
        section = doc.sections[0]
        header_p = section.header.paragraphs[0]
        header_p.text = context['title']
        header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        
        footer_p = section.footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer_p.add_run(context['page_title'] + " ") 
        self._add_page_number(run)

        doc.add_heading(f"{context['app_name']}", 0)
        doc.add_heading(f"{context['guide_type']}", 1)
        p = doc.add_paragraph()
        p.add_run(f"{context['lbl_generated']}: {context['date']}\n")
        p.add_run(f"{context['lbl_language']}: {context['language']}\n")
        p.add_run(f"{context['lbl_theme']}: {context['theme']}\n")
        p.add_run(f"{context['lbl_device']}: {context['device']}\n")
        p.add_run(f"{context['lbl_version']}: {context['version']}\n")
        p.add_run(f"{context['lbl_generated_by']}: WealthFlow Documentation Engine")
        doc.add_page_break()
        
        doc.add_heading(context['toc_title'], level=1)
        for i, item in enumerate(flat_model, 1):
            indent_level = len(item.hierarchical_number.split('.')) - 1
            p_toc = doc.add_paragraph(f"{item.hierarchical_number} {item.title}")
            p_toc.paragraph_format.left_indent = Inches(0.3 * indent_level)
        doc.add_page_break()
        
        for i, item in enumerate(flat_model, 1):
            h_level = min(1 + len(item.hierarchical_number.split('.')) - 1, 3)
            doc.add_heading(f"{item.hierarchical_number} {item.title}", level=h_level)
            
            nav_path = " > ".join(item.navigation)
            p = doc.add_paragraph()
            p.add_run(f"{context['nav_title']}: ").bold = True
            p.add_run(nav_path)
            
            if item.purpose:
                p2 = doc.add_paragraph()
                p2.add_run(f"{context['purpose_title']}: ").bold = True
                p2.add_run(item.purpose)
            
            if context['is_technical']:
                doc.add_paragraph(context['tech_notes_title'] + ":", style='Heading 3')
                doc.add_paragraph(f"Route Hierarchy: {nav_path}", style='List Bullet')
                doc.add_paragraph(f"Base Route: {getattr(item, 'route', 'N/A')}", style='List Bullet')
            elif item.steps:
                doc.add_paragraph(context['steps_title'] + ":", style='Heading 3')
                for s_idx, step in enumerate(item.steps, 1):
                    doc.add_paragraph(f"{s_idx}. {step}", style='List Paragraph')
                
            for s_idx, s_path in enumerate(item.screenshots):
                if os.path.exists(s_path):
                    try:
                        p_img = doc.add_paragraph()
                        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p_img.add_run().add_picture(s_path, width=Inches(IMAGE_DOCX_WIDTH_INCHES))
                        
                        p_cap = doc.add_paragraph()
                        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        r_cap = p_cap.add_run(f"{context['figure_title']} {i}.{s_idx+1}: {item.title}")
                        r_cap.italic = True
                    except Exception as e:
                        logger.warning(f"Could not add image {s_path} to docx: {e}")
                        
            if item.siblings:
                doc.add_paragraph("Related Pages:", style='Heading 3')
                for sib in item.siblings:
                    doc.add_paragraph(sib, style='List Bullet')
        doc.save(out_path)
