"""Recorder + runner + export through the REAL chat endpoint, the real OllamaProvider and the fake Ollama server (the model is scripted,
the pipeline is not)."""

from __future__ import annotations

import json
import os
import tempfile
from unittest.mock import patch

from django.test import TestCase
from openpyxl import load_workbook

from core.integrations.ai_provider.ollama_provider import OllamaProvider
from ai_training.models import AgentTrace
from core.models import AIConversation, AppSettings
from core.services.ai.investigation import scope
from ai_training.services import export, generator, inspect_file, runner
from core.tests.billing.test_support import grant_ai_workspace_access
from tests.support import fake_ollama as fo

from core.tests.ai.investigation.support import make_user, seed_tenant

VIEW = "core.views.ai_chat.ai_chat_core_views"
CTX = ([{"role": "system", "content": "SYS"}, {"role": "user", "content": "q"}], [])


class TraceBase(TestCase):
    def setUp(self):
        self.user = make_user("trace_user")
        seed_tenant(self.user, "Q", expenses=3, amount="50.00")
        other = make_user("trace_other")
        seed_tenant(other, "Z", expenses=7, amount="99.00")
        grant_ai_workspace_access(self.user)
        for key, value in (("ai_enabled", "true"), ("ai_agent_enabled", "true")):
            AppSettings.set(key, value)
        self.case = next(c for c in generator.generate(self.user) if c.key == "Expense:count")
        self.answers = []

    def script(self, final):
        """Model: first call looks at the data, then answers `final(n_attempt)`. Tool and answer turns are told apart by the history."""
        def handler(payload):
            if any(m.get("role") == "tool" for m in payload["messages"]):
                self.answers.append(final(len(self.answers) + 1))
                return fo.reply(self.answers[-1])
            return fo.reply(tool_calls=[fo.tool_call("app_data_query", model="Expense", aggregate={"count": "*"})])
        server = fo.FakeOllama(handler=handler).start()
        self.addCleanup(server.stop)
        return OllamaProvider(base_url=server.url, model=fo.MODEL, timeout=60)

    def run_case(self, provider):
        with patch(f"{VIEW}.get_active_ai_provider", return_value=provider), patch(f"{VIEW}.generation_pipeline.build_context", return_value=CTX):
            return runner.run(self.user, [self.case], log=lambda line: None)


class RunnerTest(TraceBase):
    def test_a_correct_repeatable_run_is_kept_and_exported(self):
        summary = self.run_case(self.script(lambda n: "You have 3 expense records."))
        self.assertEqual((summary["kept"], summary["dropped"]), (1, 0))
        traces = list(AgentTrace.objects.order_by("attempt"))
        self.assertEqual([(t.attempt, t.verdict, t.language) for t in traces], [(1, "kept", "LATIN"), (2, "kept", "LATIN")])
        first = traces[0]
        self.assertEqual([s["tool"] for s in first.steps], ["app_data_query"])
        self.assertIn('"count":3', first.steps[0]["result"])
        self.assertEqual((first.owner, first.expected, first.model_name), (self.user, {"values": [3], "meta": self.case.expected["meta"]}, fo.MODEL))
        self.assertIn("app_data_query", first.system_prompt)
        self.assertEqual(AIConversation.objects.filter(user=self.user).count(), 0)   # the runner leaves no chat history behind

    def test_export_has_only_kept_traces_in_chat_tool_format(self):
        self.run_case(self.script(lambda n: "You have 3 expense records."))
        path = os.path.join(tempfile.mkdtemp(), "t.jsonl")
        self.assertEqual(export.export_jsonl(path), 1)             # two kept runs of one question -> one line
        line = json.loads(open(path, encoding="utf-8").read().splitlines()[0])
        roles = [m["role"] for m in line["messages"]]
        self.assertEqual(roles, ["system", "user", "assistant", "tool", "assistant"])
        self.assertEqual(line["messages"][2]["tool_calls"][0]["function"]["name"], "app_data_query")
        self.assertEqual(line["messages"][-1]["content"], "You have 3 expense records.")
        self.assertTrue(line["messages"][1]["content"].startswith(self.case.question))
        self.assertTrue(line["tools"])

    def test_a_wrong_number_is_dropped_and_not_exported(self):
        summary = self.run_case(self.script(lambda n: "You have 3 expense records and 4 more."))
        self.assertEqual(summary["dropped"], 1)
        reasons = AgentTrace.objects.first().drop_reasons
        self.assertTrue(any(r.startswith("figure_not_in_tool_results") for r in reasons), reasons)
        self.assertEqual(export.export_jsonl(os.path.join(tempfile.mkdtemp(), "t.jsonl")), 0)
        rep = export.report()
        self.assertEqual((rep["questions"], rep["kept"], rep["dropped"]), (1, 0, 1))
        self.assertIn("figure_not_in_tool_results", rep["drop_reasons"])

    def test_the_right_tool_result_but_the_wrong_expected_figure_is_dropped(self):
        self.case.expected = {"values": [4]}
        self.run_case(self.script(lambda n: "You have 3 expense records."))
        self.assertIn("wrong_or_missing_figure:4", AgentTrace.objects.first().drop_reasons)

    def test_extra_figures_that_differ_between_the_two_runs_do_not_matter(self):
        calls = []

        def handler(payload):
            if any(m.get("role") == "tool" for m in payload["messages"]):
                return fo.reply("You have 3 expense records." if len(calls) == 1 else "You have 3 expense records of 50.00 each.")
            calls.append(1)
            args = {"aggregate": {"count": "*"}} if len(calls) == 1 else {"fields": ["amount"], "limit": 3}
            return fo.reply(tool_calls=[fo.tool_call("app_data_query", model="Expense", **args)])
        server = fo.FakeOllama(handler=handler).start()
        self.addCleanup(server.stop)
        self.run_case(OllamaProvider(base_url=server.url, model=fo.MODEL, timeout=60))
        self.assertEqual({tuple(t.drop_reasons) for t in AgentTrace.objects.all()}, {()})

    def test_answering_without_looking_is_dropped(self):
        server = fo.FakeOllama(handler=lambda payload: fo.reply("You have 3 expense records.")).start()
        self.addCleanup(server.stop)
        self.run_case(OllamaProvider(base_url=server.url, model=fo.MODEL, timeout=60))
        self.assertTrue(any("did_not_look_at_data" in (t.drop_reasons or []) for t in AgentTrace.objects.all()))

    def test_no_model_answer_means_no_trace_and_a_plain_reason(self):
        bad = patch(f"{VIEW}.get_active_ai_provider", return_value=None)
        with bad:
            summary = runner.run(self.user, [self.case], log=lambda line: None)
        self.assertEqual((summary["kept"], summary["dropped"], AgentTrace.objects.count()), (0, 1, 0))
        self.assertTrue(any(k.startswith("no trace") for k in summary["reasons"]), summary["reasons"])


class RecordingSwitchTest(TraceBase):
    def chat(self, provider):
        self.client.force_login(self.user)
        body = {"message": self.case.question, "chat_mode": "agent", "question_domain": "business"}
        with patch(f"{VIEW}.get_active_ai_provider", return_value=provider), patch(f"{VIEW}.generation_pipeline.build_context", return_value=CTX):
            return self.client.post(runner.CHAT_URL, json.dumps(body), content_type="application/json")

    def test_a_real_chat_is_recorded_only_when_the_setting_is_on(self):
        provider = self.script(lambda n: "You have 3 expense records.")
        self.chat(provider)
        self.assertEqual(AgentTrace.objects.count(), 0)
        AppSettings.set("ai_trace_recording", "true", user=None)
        self.chat(provider)
        trace = AgentTrace.objects.get()
        self.assertEqual((trace.source, trace.verdict, trace.expected), ("chat", "pending", None))   # judged later; no expected = not code-provable

    def test_the_ai_itself_can_never_read_the_traces(self):
        self.assertNotIn("AgentTrace", scope.catalog())


class GeneratorAndFilesTest(TraceBase):
    def more_expenses(self, n=2):
        from core.models import Expense
        base = Expense.objects.filter(owner=self.user).order_by("date").first()
        for i in range(n):
            base.pk, base.date, base.day = None, base.date.replace(day=10 + i), 10 + i
            base.save()

    def test_questions_come_from_the_users_own_rows_only(self):
        self.more_expenses()
        cases = {c.key: c for c in generator.generate(self.user)}
        self.assertEqual(cases["Expense:count"].expected["values"], [5])
        self.assertEqual(cases["Expense:count"].question, "How many expense records do I have?")
        per_cur = cases["Expense:sum_per_currency:amount"]
        self.assertEqual(per_cur.expected["values"], ["250.00"])                      # the other user's 7 x 99.00 never counts
        self.assertEqual(cases["Expense:last3:amount:date"].expected["values"], ["50.00"] * 3)

    def test_a_tie_at_the_cut_off_is_skipped(self):
        # every seeded expense has its own date, so the "last 3" question exists; two rows on the same date at the cut-off remove it
        from core.models import Expense
        self.more_expenses()
        self.assertIn("Expense:last3:amount:date", {c.key for c in generator.generate(self.user)})
        newest = Expense.objects.filter(owner=self.user).order_by("-date", "-pk")
        extra = Expense.objects.filter(owner=self.user).order_by("date").first()
        extra.pk, extra.date = None, newest[2].date
        extra.save()
        keys = {c.key for c in generator.generate(self.user)}
        self.assertFalse(any(k.startswith("Expense:last3") for k in keys))

    def test_the_same_seed_gives_the_same_sample(self):
        a = [c.key for c in generator.generate(self.user, limit=2, seed=5)]
        self.assertEqual(a, [c.key for c in generator.generate(self.user, limit=2, seed=5)])

    def test_excel_file_lists_every_run_with_its_verdict(self):
        self.case.expected = {"values": [3], "meta": self.case.expected["meta"]}
        with patch(f"{VIEW}.get_active_ai_provider", return_value=self.script(lambda n: "You have 3 expense records.")), \
                patch(f"{VIEW}.generation_pipeline.build_context", return_value=CTX):
            runner.run(self.user, [self.case], log=lambda line: None)
        path = os.path.join(tempfile.mkdtemp(), "runs.xlsx")
        self.assertEqual(inspect_file.export_xlsx(path), 2)
        rows = list(load_workbook(path).active.iter_rows(values_only=True))
        self.assertEqual(rows[0][3], "Verdict")
        self.assertEqual({r[3] for r in rows[1:]}, {"kept"})

