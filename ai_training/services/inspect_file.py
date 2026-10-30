"""The file Ehab opens on his PC: one Excel row per recorded run, with the verdict and the plain drop reason."""

from __future__ import annotations

import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from ai_training.models import AgentTrace

COLUMNS = ("Run", "Question key", "Attempt", "Verdict", "Why dropped", "Question", "Expected figures", "Answer", "Tools used (in order)",
           "Tool results (first call)", "Seconds", "Model", "Recorded at")


def export_xlsx(path: str, owner=None) -> int:
    wb = Workbook()
    ws = wb.active
    ws.title = "Agent runs"
    ws.append(list(COLUMNS))
    qs = AgentTrace.objects.exclude(source=AgentTrace.SOURCE_EVAL)
    qs = qs.filter(owner=owner) if owner else qs
    count = 0
    for t in qs.order_by("case_key", "attempt", "id"):
        ws.append([t.pk, t.case_key, t.attempt, t.verdict, "; ".join(t.drop_reasons or []), t.question,
                   json.dumps((t.expected or {}).get("values", []), ensure_ascii=False), t.answer,
                   " > ".join(s["tool"] for s in t.steps), (t.steps[0].get("result", "") if t.steps else "")[:3000], t.elapsed_s, t.model_name,
                   t.created_at.strftime("%Y-%m-%d %H:%M")])
        count += 1
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col, width in zip("ABCDEFGHIJKLM", (7, 34, 8, 10, 34, 50, 22, 60, 36, 60, 9, 18, 17)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(path)
    return count
