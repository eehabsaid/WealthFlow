"""Scores, verdict and report files of a model comparison (original / current best / new). The verdict is a recommendation only: nothing is switched.
The text report holds counts and percentages only (no figures of your data); the Excel file has the answers and stays in the git-ignored output folder."""

from __future__ import annotations

import json
import os
from collections import Counter
from statistics import median

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from .taxonomy import TITLES

MIN_QUESTIONS = 10
TOOL_FLAGS = ("did_not_look_at_data", "used_non_data_tool", "unscoped_model", "tool_error_in_run", "used_non_data_tool_or_error")
HALLUCINATION = ("figure_not_in_tool_results", "guessed_instead_of_asking")
LANGUAGE = ("answer_in_other_language", "language_not_confirmed")


def _has(row, flags) -> bool:
    return any(r.split(":")[0] in flags for r in row["reasons"])


def _pct(part, whole) -> float:
    return round(100 * part / whole, 1) if whole else 0.0


def stats(rows: list[dict]) -> dict:
    """rows = one dict per question: passed, seconds, reasons, category, lang. The warm-up is not in rows."""
    n, times = len(rows), [r["seconds"] for r in rows]
    cats = {}
    for r in rows:
        c = cats.setdefault(r.get("category", ""), [0, 0])
        c[0] += 1
        c[1] += bool(r["passed"])
    return {"questions": n, "passed": sum(1 for r in rows if r["passed"]), "rate": _pct(sum(1 for r in rows if r["passed"]), n),
            "median_s": round(median(times), 1) if times else 0.0, "mean_s": round(sum(times) / len(times), 1) if times else 0.0,
            "reasons": dict(Counter(x.split(":")[0] for r in rows for x in r["reasons"]).most_common()),
            "tool_correct": _pct(sum(1 for r in rows if not _has(r, TOOL_FLAGS)), n),
            "hallucinated": sum(1 for r in rows if _has(r, HALLUCINATION)),
            "language_ok": _pct(sum(1 for r in rows if not _has(r, LANGUAGE)), n),
            "by_category": {c: {"questions": v[0], "passed": v[1], "rate": _pct(v[1], v[0])} for c, v in cats.items()}}


def regressions(best_rows: list[dict], new_rows: list[dict]) -> list[str]:
    """Keys the current best answered right and the new model answers wrong (same question, same order)."""
    return [a["key"] for a, b in zip(best_rows, new_rows) if a["passed"] and not b["passed"]]


def verdict(old: dict, new: dict, min_gain: int, max_slower: int) -> tuple[str, str]:
    """(short verdict, one plain sentence why). `old` is the model to beat. New wins only if >= min_gain points higher AND not > max_slower % slower (median)."""
    if min(old["questions"], new["questions"]) < MIN_QUESTIONS:
        return "NOT ENOUGH QUESTIONS", f"Fewer than {MIN_QUESTIONS} test questions could be built from your data, so no fair decision is possible."
    gain = round(new["rate"] - old["rate"], 1)
    slower = round(100 * (new["median_s"] / old["median_s"] - 1), 1) if old["median_s"] else 0.0
    facts = f"correct answers {old['rate']}% -> {new['rate']}% ({gain:+} points); median time {old['median_s']}s -> {new['median_s']}s ({slower:+}%)"
    if gain >= min_gain and slower <= max_slower:
        return "NEW MODEL WINS", facts
    if gain >= min_gain:
        return "NEW MODEL IS MORE CORRECT BUT TOO SLOW", facts + f" - more than {max_slower}% slower, so not a win by your rule"
    if gain <= -min_gain:
        return "OLD MODEL IS BETTER", facts
    return "NO REAL DIFFERENCE", facts + f" - less than {min_gain} points apart"


def format_text(info: dict) -> str:
    sides = info["models"]
    head = [f"MODEL TEST   {info['when']}", "Models: " + "   ".join(f"{s['role']}: {s['name']}" for s in sides),
            f"Frozen test questions per model: {sides[0]['stats']['questions']} (same for all, never used in training; retired: {len(info['retired'])})",
            f"Ollama before each model: {info['reset']}; one warm-up question per model, not counted.", "",
            f"VERDICT ({sides[-1]['name']} against {info['rival']}): {info['verdict']}", info["why"],
            "Nothing was switched. Your model in Settings is unchanged; you decide.", "", f"{'':34}" + "".join(f"{s['role'].upper():>14}" for s in sides)]
    rows = [("Correct answers", "passed"), ("Correct %", "rate"), ("Tool use correct %", "tool_correct"), ("Hallucinated answers", "hallucinated"),
            ("Language consistent %", "language_ok"), ("Median seconds", "median_s")]
    body = [f"{label:34}" + "".join(f"{s['stats'][key]:>14}" for s in sides) for label, key in rows]
    cats = sorted({c for s in sides for c in s["stats"]["by_category"]})
    body += ["", "ACCURACY BY CATEGORY (correct / asked)"]
    for c in cats:
        cells = "".join(f"{(lambda v: f'{v['passed']}/{v['questions']}' if v else '-')(s['stats']['by_category'].get(c)):>14}" for s in sides)
        body.append(f"{TITLES.get(c, c or 'other')[:34]:34}{cells}")
    body += ["", "REGRESSIONS (the current best was right, the new model is wrong): " + str(len(info["regressions"]))]
    body += [f"  {k.split('|')[0]}" for k in info["regressions"][:40]]
    for s in sides:
        if s["stats"]["reasons"]:
            body.append(f"{s['role']} failed because: " + ", ".join(f"{k} x{v}" for k, v in s["stats"]["reasons"].items()))
    body += ["", "Language check: Arabic by script, French/German/English by the app's own translation words; an answer that cannot be confirmed in French or German counts as wrong.",
             "Categories code cannot fully prove (what-if reasoning, explanation quality, anomalies, ambiguity) are judged only on the provable part."]
    return "\n".join(head + body)


def write_files(out_dir: str, info: dict) -> dict:
    os.makedirs(out_dir, exist_ok=True)
    paths = {k: os.path.join(out_dir, f) for k, f in (("report", "eval_report.txt"), ("json", "eval_results.json"), ("excel", "eval_runs.xlsx"))}
    with open(paths["report"], "w", encoding="utf-8") as fh:
        fh.write(format_text(info) + "\n")
    with open(paths["json"], "w", encoding="utf-8") as fh:
        json.dump(info, fh, ensure_ascii=False, indent=2, default=str)
    wb = Workbook()
    ws = wb.active
    ws.title = "Model test"
    sides = info["models"]
    ws.append(["Question key", "Category", "Language", "Question", "Expected figures"] + [f"{s['role']}: {h}" for s in sides for h in ("correct", "seconds", "why wrong", "answer")])
    for i, first in enumerate(sides[0]["rows"]):
        line = [first["key"], first.get("category", ""), first.get("lang", ""), first["question"], json.dumps(first["expected"].get("values"), ensure_ascii=False)]
        for s in sides:
            r = s["rows"][i]
            line += ["yes" if r["passed"] else "NO", r["seconds"], "; ".join(r["reasons"]), r["answer"]]
        ws.append(line)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(paths["excel"])
    return paths
