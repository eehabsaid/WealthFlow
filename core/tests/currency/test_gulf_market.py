"""Gulf market (SAR/AED base): no EGP anywhere, spot gold, safe catalog switch."""
import io
from decimal import Decimal
from types import SimpleNamespace

from django.contrib.auth import get_user_model
from django.test import TestCase
from openpyxl import load_workbook

from core.models import BalanceEntry, Currency, ExchangeRate, GoldPrice
from core.reports.excel_sheets_builder.build_balance_sheet.context import resolve_slot_codes
from core.services.fixed_assets.gold_sync_service.sync import _refresh_gold_asset_pricing
from core.services.shared.market_profile import is_gulf_user, spot_gold_snapshot

User = get_user_model()

RATES = {"USD": "52.30", "EUR": "59.05", "GBP": "69.12", "SAR": "13.94", "AED": "14.23", "EGP": "1"}


def _codes(response):
    return {row["code"].upper() for row in response.json()["currencies"]}


class GulfMarketBase(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="gulf_user", password="pw-12345-x")
        self.client.force_login(self.user)
        for code, value in RATES.items():
            if code == "EGP":
                continue
            ExchangeRate.objects.create(
                currency_code=code, currency_name=code, buy_rate=Decimal(value),
                sell_rate=Decimal(value), mid_rate=Decimal(value),
            )
        self.gold = GoldPrice.objects.create(
            carat_24k=7000, carat_22k=6400, carat_21k=6125, carat_18k=5250,
            carat_24k_buy=6900, carat_22k_buy=6300, carat_21k_buy=6025, carat_18k_buy=5150,
            usd_gram_24k=Decimal("134.000000"), usd_per_oz=Decimal("4168.0000"),
            usd_to_egp=Decimal("52.300000"),
        )

    def switch(self, code):
        return self.client.post("/api/base-currency/", data={"code": code}, content_type="application/json")


class CatalogSwitchTests(GulfMarketBase):
    def test_new_user_can_pick_aed_in_the_wizard(self):
        self.assertIn("AED", _codes(self.client.get("/api/currencies/")))
        self.assertEqual(self.client.post("/api/onboarding/complete/", data={"default_currency": "AED"},
                                          content_type="application/json").status_code, 200)
        self.assertEqual(self.client.get("/api/base-currency/").json()["code"], "AED")
        self.assertNotIn("EGP", _codes(self.client.get("/api/currencies/")))

    def test_spot_gold_is_rounded_to_two_decimals(self):
        self.switch("SAR")
        gold = self.client.get("/api/gold/").json()["gold"]
        self.assertEqual(gold["carat_24k"], round(gold["carat_24k"], 2))

    def test_new_user_starts_with_egp(self):
        self.assertIn("EGP", _codes(self.client.get("/api/currencies/")))

    def test_switch_to_aed_builds_gulf_catalog_without_egp(self):
        self.assertEqual(self.switch("AED").status_code, 200)
        codes = _codes(self.client.get("/api/currencies/"))
        self.assertEqual(codes, {"SAR", "AED", "USD", "EUR", "GOLD"})
        self.assertTrue(is_gulf_user(self.user))

    def test_switch_to_sar_has_no_egp(self):
        self.assertEqual(self.switch("SAR").status_code, 200)
        self.assertNotIn("EGP", _codes(self.client.get("/api/currencies/")))

    def test_switch_back_to_egp_restores_the_row(self):
        self.switch("SAR")
        self.assertEqual(self.switch("EGP").status_code, 200)
        self.assertIn("EGP", _codes(self.client.get("/api/currencies/")))
        self.assertFalse(is_gulf_user(self.user))

    def test_used_egp_is_kept_and_its_data_survives(self):
        egp = Currency.objects.get(owner=self.user, code="EGP")
        entry = BalanceEntry.objects.create(owner=self.user, title="Cairo", currency=egp, amount=500)
        self.assertEqual(self.switch("SAR").status_code, 200)
        self.assertTrue(Currency.objects.filter(owner=self.user, code="EGP").exists())
        self.assertTrue(BalanceEntry.objects.filter(pk=entry.pk).exists())

    def test_failed_switch_does_not_touch_the_catalog(self):
        ExchangeRate.objects.filter(currency_code="AED").delete()
        self.assertEqual(self.switch("AED").status_code, 400)
        self.assertIn("EGP", _codes(self.client.get("/api/currencies/")))

    def test_other_users_catalog_is_untouched(self):
        other = User.objects.create_user(username="egy_user", password="pw-12345-x")
        self.switch("AED")
        self.assertTrue(Currency.objects.filter(owner=other, code="EGP").exists())


class RatesAndGoldTests(GulfMarketBase):
    def setUp(self):
        super().setUp()
        ExchangeRate.objects.create(
            currency_code="EGP", currency_name="Egyptian Pound", buy_rate=Decimal("0.019"),
            sell_rate=Decimal("0.019"), mid_rate=Decimal("0.019"),
        )

    def _rate_codes(self):
        return {r["currency_code"] for r in self.client.get("/api/rates/").json()["rates"]}

    def test_egyptian_user_still_sees_egp_rate(self):
        self.assertIn("EGP", self._rate_codes())

    def test_gulf_user_does_not_see_egp_rate_but_sees_aed(self):
        self.switch("SAR")
        codes = self._rate_codes()
        self.assertNotIn("EGP", codes)
        self.assertIn("AED", codes)

    def test_gulf_gold_is_spot_in_base_currency(self):
        self.switch("SAR")
        gold = self.client.get("/api/gold/").json()["gold"]
        self.assertEqual(gold["market"], "spot")
        self.assertEqual(gold["currency"], "SAR")
        expected_rate = float(Decimal(RATES["USD"]) / Decimal(RATES["SAR"]))  # SAR per 1 USD
        self.assertAlmostEqual(gold["usd_to_base"], expected_rate, places=3)
        self.assertAlmostEqual(gold["carat_24k"], 134 * expected_rate, places=1)
        self.assertAlmostEqual(gold["carat_22k"], 134 * expected_rate * 22 / 24, places=1)
        self.assertEqual(gold["carat_24k_buy"], gold["carat_24k"])

    def test_egyptian_gold_is_unchanged_dealer_price(self):
        gold = self.client.get("/api/gold/").json()["gold"]
        self.assertNotIn("market", gold)
        self.assertEqual(gold["currency"], "EGP")
        self.assertEqual(gold["carat_24k"], 7000)

    def test_gold_refresh_view_also_returns_user_scoped_prices(self):
        self.switch("AED")
        from core.views.settings.market import gold_price_views as views
        data = views._gold_for_user(self.user, self.gold)
        self.assertEqual(data["currency"], "AED")

    def test_asset_sync_values_gulf_gold_in_base_not_egp(self):
        self.switch("SAR")
        snap = spot_gold_snapshot(self.user, self.gold)
        details = SimpleNamespace(
            purity="21k", unit="gram", weight=Decimal("10"), market_price=0, cashback_per_gram=0,
            save=lambda **kw: None,
        )
        asset = SimpleNamespace(
            owner=self.user, asset_type="Gold", purchase_price=Decimal("3750"),
            purchase_usd_rate=0, purchase_price_usd=0, current_market_value=0,
            valuation_source="", last_valuation_date=None, save=lambda **kw: None,
        )
        _refresh_gold_asset_pricing(asset, gold_details=details, latest_gold_price=self.gold)
        self.assertAlmostEqual(float(asset.current_market_value), snap["carat_21k"] * 10, places=1)
        self.assertAlmostEqual(float(asset.purchase_usd_rate), snap["usd_to_base"], places=4)
        self.assertAlmostEqual(float(asset.purchase_price_usd), 3750 / snap["usd_to_base"], places=2)


class BalanceSheetSlotTests(TestCase):
    def test_egp_base_layout_is_unchanged(self):
        self.assertEqual(resolve_slot_codes("EGP"), ("USD", "EUR", "SAR"))

    def test_sar_base_replaces_the_duplicate_sar_column(self):
        self.assertEqual(resolve_slot_codes("SAR", gulf=True), ("USD", "EUR", "AED"))

    def test_usd_base_replaces_the_duplicate_usd_column(self):
        slots = resolve_slot_codes("USD")
        self.assertNotIn("USD", slots)
        self.assertEqual(len(set(slots)), 3)

    def test_eur_base_and_aed_base(self):
        self.assertNotIn("EUR", resolve_slot_codes("EUR"))
        self.assertEqual(resolve_slot_codes("AED", gulf=True), ("USD", "EUR", "SAR"))


class ExcelExportGulfTests(GulfMarketBase):
    def _workbook(self):
        response = self.client.get("/api/export/excel/")
        self.assertEqual(response.status_code, 200, response.content[:300])
        return load_workbook(io.BytesIO(b"".join(response.streaming_content) if response.streaming else response.content))

    def _texts(self, wb):
        out = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        out.append(str(cell.value))
                    out.append(cell.number_format)
        return out

    def _seed_balances(self, home_codes):
        for code in home_codes:
            cur = Currency.objects.get(owner=self.user, code=code)
            BalanceEntry.objects.create(owner=self.user, title="Home Balance", currency=cur, amount=100)

    def test_sar_user_workbook_has_no_egp_and_no_duplicate_columns(self):
        self.switch("SAR")
        self._seed_balances(["SAR", "USD"])
        wb = self._workbook()
        bad = [t for t in self._texts(wb) if "EGP" in t or "ج.م" in t or "الجنيه الذهب" in t]
        self.assertEqual(bad, [])
        header = [c.value for c in wb["BALANCE"][1]][:6]
        self.assertEqual(header[1:5], ["SAR", "USD", "EUR", "AED"])
        total = str(wb["BALANCE"].cell(row=wb["BALANCE"].max_row, column=2).value)
        formulas = [str(c.value) for row in wb["BALANCE"].iter_rows() for c in row if str(c.value).startswith("=B")]
        self.assertTrue(formulas)
        self.assertTrue(all("'Exchange Rates'!" not in f for f in formulas), total)

    def test_usd_user_has_no_duplicate_usd_column(self):
        self.switch("USD")
        wb = self._workbook()
        header = [c.value for c in wb["BALANCE"][1]][1:5]
        self.assertEqual(header[0], "USD")
        self.assertEqual(header.count("USD"), 1)

    def test_egyptian_user_layout_unchanged(self):
        wb = self._workbook()
        header = [c.value for c in wb["BALANCE"][1]][:6]
        self.assertEqual(header, ["Title", "EGP", "USD", "EUR", "SAR", "Gold"])
