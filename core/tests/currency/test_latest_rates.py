"""One rule for "the current rate of a currency": newest fetched_at, ties -> higher id (same as CurrencyConversionService)."""

from __future__ import annotations

import datetime
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.test import TestCase

from core.models import ExchangeRate
from core.services.ai.providers.market_provider import MarketDataProvider
from core.services.ai.query_engine.executors.market import _latest_rows
from core.services.fixed_assets.usd_rate_service import UsdRateService
from core.services.shared.latest_rates import latest_exchange_rates

NOW = datetime.datetime(2026, 10, 1, 12, 0, tzinfo=datetime.timezone.utc)


def rate(code, buy, when):
    row = ExchangeRate.objects.create(currency_code=code, buy_rate=Decimal(buy), sell_rate=Decimal(buy), mid_rate=Decimal(buy))
    ExchangeRate.objects.filter(pk=row.pk).update(fetched_at=when)      # fetched_at is auto_now_add: back-fill it explicitly
    return row


class LatestRatesTest(TestCase):
    def setUp(self):
        # USD: the higher-id row is an OLDER back-filled fetch; the lower-id row is the genuinely newer one
        self.usd_new = rate("USD", "50.00", NOW)
        self.usd_backfill = rate("USD", "40.00", NOW - datetime.timedelta(days=30))
        # SAR: two rows with the same timestamp -> the higher id wins
        self.sar_a = rate("SAR", "13.00", NOW)
        self.sar_b = rate("SAR", "13.50", NOW)
        self.eur = rate("EUR", "55.00", NOW)

    def test_helper_returns_one_current_row_per_currency(self):
        ids = set(latest_exchange_rates().values_list("id", flat=True))
        self.assertEqual(ids, {self.usd_new.id, self.sar_b.id, self.eur.id})

    def test_every_screen_that_shows_the_current_rate_agrees(self):
        provider = {r["currency_code"]: r for r in MarketDataProvider().get_data(None)["exchange_rates"]}
        self.assertEqual(Decimal(provider["USD"]["buy_rate"]), Decimal("50.00"))
        self.assertEqual(Decimal(provider["SAR"]["buy_rate"]), Decimal("13.50"))
        rows = _latest_rows()
        self.assertEqual(rows["USD"].id, self.usd_new.id)
        self.assertEqual(UsdRateService()._latest_rate_by_code()["USD"].id, self.usd_new.id)
        user = get_user_model().objects.create_user("lr_user", password="x")
        self.client.force_login(user)
        listed = {r["currency_code"]: r for r in self.client.get("/api/rates/").json()["rates"]}
        self.assertEqual(float(listed["USD"]["buy_rate"]), 50.0)
        self.assertEqual(float(listed["SAR"]["buy_rate"]), 13.5)

    def test_it_matches_the_conversion_service_rule(self):
        from core.services.shared.currency_conversion_service import CurrencyConversionService

        self.assertEqual(CurrencyConversionService.get_latest_buy_rate("USD"), Decimal("50.00"))
        self.assertEqual(latest_exchange_rates().get(currency_code="USD").buy_rate, Decimal("50.00"))

    def test_excel_export_lists_the_same_current_rows(self):
        from openpyxl import load_workbook

        from core.reports.excel_main_generator import generate_excel

        user = get_user_model().objects.create_user("lr_excel", password="x")
        buf = generate_excel(user, lang="en")
        buf.seek(0)
        sheet = load_workbook(buf)["Exchange Rates"]
        values = [c.value for row in sheet.iter_rows() for c in row if c.value is not None]
        numbers = {float(v) for v in values if isinstance(v, (int, float)) or str(v).replace(".", "", 1).isdigit()}
        self.assertIn(50.0, numbers)         # USD current rate
        self.assertNotIn(40.0, numbers)      # the older back-filled USD row is not listed
        self.assertIn(13.5, numbers)         # SAR tie -> higher id
        self.assertNotIn(13.0, numbers)
