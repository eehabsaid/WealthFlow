"""Excel 'Exchange Rates' sheet must be quoted in the report owner's base
currency even when the stored rates are pivoted on another currency (Gulf
market pivot). Regression: an EGP user's workbook converted USD/EUR/SAR with
AED-quoted cells, so foreign-currency totals were never calculated in EGP."""
from decimal import Decimal

from django.test import TestCase
from openpyxl import Workbook

from core.models import ExchangeRate
from core.reports.excel_sheets_builder.exchange_rates import build_exchange_rates_sheet
from core.reports.report_context import set_report_base_code
from core.services.shared.currency_conversion_service import set_rate_pivot_code


def _rate(code, buy):
    return ExchangeRate.objects.create(
        currency_code=code, buy_rate=Decimal(buy), sell_rate=Decimal(buy), mid_rate=Decimal(buy)
    )


def _sheet(rates):
    ws = Workbook().active
    build_exchange_rates_sheet(ws, rates, [], owner=None)
    return ws


class ExchangeRatesSheetBaseTests(TestCase):
    def tearDown(self):
        set_report_base_code("EGP")
        set_rate_pivot_code("EGP")

    def test_aed_pivot_rates_are_rebased_to_egp(self):
        set_rate_pivot_code("AED")
        set_report_base_code("EGP")
        rates = [_rate("EGP", "0.0750"), _rate("USD", "3.6500"), _rate("SAR", "0.9750")]
        ws = _sheet(rates)
        self.assertAlmostEqual(ws["B2"].value, 3.65 / 0.075, places=4)  # USD row
        self.assertAlmostEqual(ws["B11"].value, 0.975 / 0.075, places=4)  # SAR row

    def test_rates_unchanged_when_pivot_is_the_base(self):
        set_rate_pivot_code("EGP")
        set_report_base_code("EGP")
        ws = _sheet([_rate("USD", "48.5"), _rate("SAR", "12.9")])
        self.assertAlmostEqual(ws["B2"].value, 48.5, places=6)
        self.assertAlmostEqual(ws["B11"].value, 12.9, places=6)
