from decimal import Decimal

from django.contrib.auth import get_user_model
from django.test import TestCase

from core.models import Bank, BalanceEntry, Currency, ExchangeRate, UserProfile
from core.reports.report_context import get_report_base_code, set_report_base_code
from core.reports.report_utils import get_text
from core.reports.excel_formatting_helpers import fmt_for_code, fmt_for_code_cert, fmt_base, fmt_base_cert
from core.reports.excel_main_generator import generate_excel

User = get_user_model()


class ReportContextTests(TestCase):
    def test_defaults_to_egp_when_never_set(self):
        # a fresh thread with nothing set yet
        import threading

        result = {}

        def check():
            result["code"] = get_report_base_code()

        t = threading.Thread(target=check)
        t.start()
        t.join()
        self.assertEqual(result["code"], "EGP")

    def test_set_and_get_round_trip(self):
        set_report_base_code("sar")
        self.assertEqual(get_report_base_code(), "SAR")
        set_report_base_code("EGP")  # reset for other tests sharing this thread


class GetTextBaseSubstitutionTests(TestCase):
    """The {base} token bug: server-side get_text() never substituted it,
    so labels like "Amount ({base})" were printed to PDFs/Excel files
    verbatim, with the literal braces, for every user."""

    def test_substitutes_base_token_from_report_context(self):
        set_report_base_code("SAR")
        try:
            text = get_text("amount_egp", "en", {"amount_egp": "Amount ({base})"})
            self.assertEqual(text, "Amount (SAR)")
        finally:
            set_report_base_code("EGP")

    def test_leaves_text_without_token_untouched(self):
        text = get_text("amount", "en", {"amount": "Amount"})
        self.assertEqual(text, "Amount")

    def test_default_used_when_key_missing_also_gets_substituted(self):
        set_report_base_code("USD")
        try:
            text = get_text("missing_key", "en", {}, default="Fallback ({base})")
            self.assertEqual(text, "Fallback (USD)")
        finally:
            set_report_base_code("EGP")


class DynamicFormatHelperTests(TestCase):
    def test_fmt_base_reads_report_context(self):
        set_report_base_code("SAR")
        try:
            self.assertIn("SAR", fmt_base())
            self.assertIn("SAR", fmt_base_cert())
        finally:
            set_report_base_code("EGP")

    def test_fmt_for_code_uses_explicit_code_not_report_context(self):
        # A bank certificate/expense has its own currency, independent of
        # the report owner's base — fmt_for_code(explicit_code) must use
        # that explicit code even when the report context says something
        # else entirely.
        set_report_base_code("SAR")
        try:
            self.assertIn("USD", fmt_for_code("USD"))
            self.assertIn("USD", fmt_for_code_cert("usd"))
        finally:
            set_report_base_code("EGP")

    def test_fmt_for_code_falls_back_to_report_context_when_no_code_given(self):
        set_report_base_code("QAR")
        try:
            self.assertIn("QAR", fmt_for_code(None))
        finally:
            set_report_base_code("EGP")


class BalanceReportViewCurrencyNeutralTests(TestCase):
    """Regression test for the real bug: total_egp used to filter
    currency__code="EGP" only, so a bank holding non-EGP balances (or any
    balance at all for a non-EGP-base user) reported an incomplete/zero
    total, even though the frontend already labels this as "total in your
    base currency"."""

    def setUp(self):
        self.user = User.objects.create_user(username="sar_report_user", password="pw12345")
        self.egp, _ = Currency.objects.get_or_create(owner=self.user, code="EGP", defaults={"symbol": "EGP", "name": "EGP"})
        self.usd, _ = Currency.objects.get_or_create(owner=self.user, code="USD", defaults={"symbol": "$", "name": "USD"})
        self.sar, _ = Currency.objects.get_or_create(owner=self.user, code="SAR", defaults={"symbol": "SAR", "name": "SAR"})
        profile, _ = UserProfile.objects.get_or_create(user=self.user)
        profile.preferred_currency = "SAR"
        profile.save(update_fields=["preferred_currency"])
        self.bank = Bank.objects.create(owner=self.user, name="Test Bank")

        # Pivot = EGP (the historical default). 1 SAR = 13 EGP, 1 USD = 50 EGP.
        ExchangeRate.objects.create(currency_code="SAR", buy_rate=Decimal("13"), sell_rate=Decimal("13"), mid_rate=Decimal("13"))
        ExchangeRate.objects.create(currency_code="USD", buy_rate=Decimal("50"), sell_rate=Decimal("50"), mid_rate=Decimal("50"))

        BalanceEntry.objects.create(
            owner=self.user, title="SAR cash", balance_type=BalanceEntry.BalanceType.CASH,
            bank=self.bank, currency=self.sar, amount=Decimal("100"),
        )
        BalanceEntry.objects.create(
            owner=self.user, title="USD cash", balance_type=BalanceEntry.BalanceType.CASH,
            bank=self.bank, currency=self.usd, amount=Decimal("10"),
        )

    def test_total_converts_every_currency_to_the_users_own_base(self):
        self.client.force_login(self.user)
        resp = self.client.get("/api/reports/balance/")
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        bank_row = next(b for b in data["by_bank"] if b["bank_id"] == self.bank.id)
        # 100 SAR (already base) + 10 USD * (50 EGP/USD / 13 EGP/SAR) = 100 + 38.46...
        expected = 100 + 10 * (50 / 13)
        self.assertAlmostEqual(bank_row["total_egp"], expected, places=2)

    def test_egp_only_filter_bug_is_gone(self):
        # Before the fix this would have been 0, since neither entry above
        # is literally in EGP.
        self.client.force_login(self.user)
        resp = self.client.get("/api/reports/balance/")
        data = resp.json()
        bank_row = next(b for b in data["by_bank"] if b["bank_id"] == self.bank.id)
        self.assertGreater(bank_row["total_egp"], 0)


class GenerateExcelCurrencyNeutralTests(TestCase):
    """End-to-end smoke test: build a real workbook for a non-EGP-base user
    and check the Balance sheet reflects that base, not a hardcoded EGP."""

    def setUp(self):
        self.user = User.objects.create_user(username="sar_excel_user", password="pw12345")
        self.sar, _ = Currency.objects.get_or_create(owner=self.user, code="SAR", defaults={"symbol": "SAR", "name": "SAR"})
        self.usd, _ = Currency.objects.get_or_create(owner=self.user, code="USD", defaults={"symbol": "$", "name": "USD"})
        profile, _ = UserProfile.objects.get_or_create(user=self.user)
        profile.preferred_currency = "SAR"
        profile.save(update_fields=["preferred_currency"])
        self.bank = Bank.objects.create(owner=self.user, name="Test Bank")
        ExchangeRate.objects.create(currency_code="SAR", buy_rate=Decimal("13"), sell_rate=Decimal("13"), mid_rate=Decimal("13"))
        ExchangeRate.objects.create(currency_code="USD", buy_rate=Decimal("50"), sell_rate=Decimal("50"), mid_rate=Decimal("50"))
        BalanceEntry.objects.create(
            owner=self.user, title="SAR cash", balance_type=BalanceEntry.BalanceType.CASH,
            bank=self.bank, currency=self.sar, amount=Decimal("250"),
        )

    def test_balance_sheet_header_uses_the_owners_own_base_currency(self):
        import openpyxl
        import io

        buf = generate_excel(self.user)
        wb = openpyxl.load_workbook(io.BytesIO(buf.getvalue()) if hasattr(buf, "getvalue") else buf)
        ws = wb["BALANCE"] if "BALANCE" in wb.sheetnames else wb.worksheets[0]
        header_row = [c.value for c in ws[1]]
        self.assertIn("SAR", header_row)
        self.assertNotIn("EGP", header_row)

    def test_generation_does_not_crash_for_egp_base_user_either(self):
        egp_user = User.objects.create_user(username="egp_excel_user", password="pw12345")
        Currency.objects.get_or_create(owner=egp_user, code="EGP", defaults={"symbol": "EGP", "name": "EGP"})
        # No profile row at all -> falls back to platform default (EGP)
        generate_excel(egp_user)  # must not raise
