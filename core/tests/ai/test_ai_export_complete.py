"""Downloads hold the COMPLETE data, not the chat's capped view; the `generate_file` tool; settings, audit, routing.

Covers what the first export version did not: rows beyond the 300-row chat table, model answers that drop rows, an AI
tool that builds a file from a database query (no 100-row cap, owner isolation), the Settings switch / limits, the audit
entry, give-up phrases that hand over to the investigator, and code files."""

import datetime
import re
from decimal import Decimal
from io import BytesIO
from unittest.mock import MagicMock, patch

from django.test import TestCase
from openpyxl import load_workbook

from core.integrations.ai_provider import OllamaProvider
from core.models import AppSettings, AuthAuditLog, Expense
from core.services.ai.exports import requested
from core.services.ai.tools import get_registered_tool_schemas
from core.tests.ai import qe_fixtures as fx
from core.tests.ai.investigation.test_chat_investigator import GP, VIEW

URL = "/api/financial-advisor/ai/messages/{id}/export/{fmt}/"
LINK = re.compile(r"/api/financial-advisor/ai/messages/(\d+)/export/(\w+)/")


def grid(response):
    return [[c.value for c in r] for r in load_workbook(BytesIO(response.content)).active.iter_rows()]


def bulk_expenses(owner, category, n, currency=None):
    Expense.objects.bulk_create([Expense(owner=owner, category=category, date=datetime.date(2026, 6, 1 + i % 25), year=2026, month=6, amount=Decimal("1"),
                                         amount_base=Decimal("1"), description=f"bulk {i}") for i in range(n)])



class FullDataTests(TestCase):
    def setUp(self):
        fx.build(self)
        self.category = Expense.objects.filter(owner=self.user).first().category

    def ask(self, text):
        return fx.ask(self, text, conversation=False)["message"]

    def test_file_has_every_row_while_the_chat_table_is_capped(self):
        bulk_expenses(self.user, self.category, 450)
        msg = self.ask("list my expenses for June 2026 and export to excel")
        self.assertIn("Showing the first 300", msg["content"])                           # the chat view is capped...
        rows = grid(self.client.get(URL.format(id=msg["id"], fmt="xlsx")))
        self.assertEqual(len([r for r in rows if r[3] and str(r[3]).startswith("bulk")]), 450)   # ...the file is not
        self.assertNotIn("Showing the first", str(rows))

    def test_instant_answers_off_still_answers_an_export_request_without_the_model(self):
        """Export is its own feature: the exact table + file come from the router with 0 LLM calls, whatever the switch says."""
        bulk_expenses(self.user, self.category, 60)
        AppSettings.set("ai_direct_answers", "false")
        with patch.object(OllamaProvider, "generate", side_effect=AssertionError("LLM must not be called")):
            data = fx.ask(self, "list my expenses for June 2026 and export it to excel", conversation=False)
        rows = grid(self.client.get(URL.format(id=data["message"]["id"], fmt="xlsx")))
        self.assertEqual(len([r for r in rows if r[3] and str(r[3]).startswith("bulk")]), 60)
        self.assertEqual(data["message"]["tool_calls"][0]["llm_calls"], 0)

    def test_follow_up_export_of_a_model_answer_is_rebuilt_from_the_database(self):
        """Direct Answers OFF: the model wrote a shortened table; 'export it to excel' still yields every row."""
        bulk_expenses(self.user, self.category, 60)
        AppSettings.set("ai_direct_answers", "false")
        prov = MagicMock()
        prov.supports_tools = False
        prov.generate.side_effect = [{"content": "| Date | Description | Amount |\n|---|---|---|\n| 2026-06-01 | bulk 0 | 1.00 EGP |", "tool_calls": None, "error": None}] * 3
        ctx = ([{"role": "system", "content": "SYS"}, {"role": "user", "content": "q"}], ["expenses"])
        with patch(f"{VIEW}.get_active_ai_provider", return_value=prov), patch(GP, return_value=ctx):
            first = fx.ask(self, "list my expenses for June 2026")
        self.assertNotIn("/export/", first["message"]["content"])
        follow = fx.ask(self, "export it to excel")["message"]
        m = LINK.search(follow["content"])
        self.assertEqual(m.group(1), str(first["message"]["id"]))
        rows = grid(self.client.get(URL.format(id=m.group(1), fmt="xlsx")))
        self.assertEqual(len([r for r in rows if r[3] and str(r[3]).startswith("bulk")]), 60)

    def test_row_limit_setting_caps_the_file_and_says_so(self):
        bulk_expenses(self.user, self.category, 450)
        AppSettings.set("ai_file_max_rows", "100")
        msg = self.ask("list my expenses for June 2026 and export to excel")
        rows = grid(self.client.get(URL.format(id=msg["id"], fmt="xlsx")))
        self.assertEqual(len([r for r in rows if r[0] == "Jun 2026"]), 100)
        self.assertIn("Showing the first 100", str(rows))

    def test_switch_off_removes_links_downloads_and_tool(self):
        msg = self.ask("list my expenses for 2026 and export to excel")
        self.assertTrue(LINK.search(msg["content"]))
        AppSettings.set("ai_file_generation", "false")
        self.assertEqual(self.client.get(URL.format(id=msg["id"], fmt="xlsx")).status_code, 404)
        self.assertFalse(LINK.search(self.ask("list my expenses for 2026 and export to excel")["content"]))
        self.assertEqual(requested("export to excel"), [])
        self.assertNotIn("generate_file", {s["function"]["name"] for s in get_registered_tool_schemas()})

    def test_size_limit_answers_413(self):
        bulk_expenses(self.user, self.category, 300)
        msg = self.ask("list my expenses for June 2026 and export to pdf")
        AppSettings.set("ai_file_max_mb", "1")
        with patch("core.views.ai_chat.ai_export_views.build_file", return_value=(b"x" * (1024 * 1024 + 1), "a.pdf")):
            self.assertEqual(self.client.get(URL.format(id=msg["id"], fmt="pdf")).status_code, 413)

    def test_download_is_audited_and_not_cached(self):
        msg = self.ask("list my expenses for 2026 and export to excel")
        r = self.client.get(URL.format(id=msg["id"], fmt="xlsx"))
        self.assertEqual(r["Cache-Control"], "private, no-store")
        entry = AuthAuditLog.objects.filter(user=self.user, event_type="ai_file_downloaded").get()
        self.assertIn(f"message {msg['id']}", entry.details)

    def test_every_capability_downloads_in_every_document_format(self):
        for q in ("list my expenses for 2026", "what is my total balance", "gold price today", "my exchange rates", "list my certificates"):
            for fmt in ("xlsx", "docx", "pdf", "csv", "txt"):
                msg = self.ask(f"{q} and export it as {fmt}")
                m = LINK.search(msg["content"])
                self.assertIsNotNone(m, (q, fmt))
                self.assertEqual(self.client.get(URL.format(id=m.group(1), fmt=m.group(2))).status_code, 200, (q, fmt))
