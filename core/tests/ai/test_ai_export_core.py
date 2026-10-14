"""Downloadable answers, core: intent words (en/ar/fr/de), markdown -> document, typed cells, the registry,
and every file format built from one sample answer (Excel, Word, PDF, CSV, TXT; English and Arabic)."""

import csv
import re
from io import BytesIO, StringIO

from django.test import SimpleTestCase
from django.urls import reverse
from docx import Document
from openpyxl import load_workbook

from core.services.ai.exports import ExportFormat, all_formats, default_formats, detect, get_format, register, strip_export_terms
from core.services.ai.query_engine.slots import route
from core.services.ai.exports.cells import parse_number, typed_table
from core.services.ai.exports.document import EXPORT_PATH_MARK, ExportDocument, parse_markdown, plain, strip_export_lines
from core.services.ai.exports.registry import _FORMATS

LINK = "[Download as Excel (.xlsx)](/api/financial-advisor/ai/messages/5/export/xlsx/) · [Download as PDF (.pdf)](/api/financial-advisor/ai/messages/5/export/pdf/)"
SAMPLE = f"""## Salary for 2026 (amounts in EGP):

| Month | Company | Expected | Paid | Bonus | Change |
|---|---|---|---|---|---|
| Jan 2026 | Acme | 90,000.00 EGP | 87,643.86 EGP | 0.00 EGP | +9.55% |
| Feb 2026 | Acme | 90,000.00 EGP | 88,000.50 EGP | 100.00 EGP | n/a |
| **Grand Total** | | **180,000.00 EGP** | **175,644.36 EGP** | **100.00 EGP** | |

- First note
1. Second step

Plain paragraph with **bold** word.

{LINK}"""
AR_SAMPLE = "| الشهر | المبلغ |\n|---|---|\n| يناير 2026 | 1,000.00 EGP |\n| **الإجمالي الكلي** | **1,000.00 EGP** |\n\nمرحبا بكم"
PHRASES = {
    "xlsx": ["list salaries 2026 and export it in excel sheet", "put this in a spreadsheet", "اعرض رواتب 2026 وصدرها في ملف اكسل", "exporte mes salaires en excel", "Gehälter als Excel herunterladen"],
    "docx": ["list my expenses and download as word", "give me this as a word document", "صدر اسعار الذهب الي ملف ورد", "exporte ça au format word", "lade das als Word herunter"],
    "pdf": ["export the gold prices to pdf", "حمل الرواتب بي دي اف", "télécharge le taux de change en pdf", "exportiere die Ausgaben als PDF"],
    "csv": ["save it as csv"], "txt": ["as a text file please"],
}
NO_FILE = ["what is my balance", "how do I export my salaries", "how can I download a pdf report", "I need a word of advice", "the word 'salary' means pay", "كيف اصدر الرواتب", ""]


def doc_from(text, lang="en"):
    return ExportDocument(title="WealthFlow AI", question="Question: q", generated="Generated on 2026-10-12", lang=lang, blocks=parse_markdown(text))


class IntentTests(SimpleTestCase):
    def test_named_formats_in_every_language(self):
        for key, phrases in PHRASES.items():
            for phrase in phrases:
                with self.subTest(phrase=phrase):
                    self.assertEqual([f.key for f in detect(phrase)], [key])

    def test_several_named_formats_keep_registry_order(self):
        self.assertEqual([f.key for f in detect("export it in pdf and word and excel")], ["xlsx", "docx", "pdf"])

    def test_bare_export_offers_the_default_formats(self):
        for phrase in ("export it", "download exchange rates", "صدر الرواتب", "exportiere das"):
            with self.subTest(phrase=phrase):
                self.assertEqual([f.key for f in detect(phrase)], [f.key for f in default_formats()])
        self.assertEqual([f.key for f in default_formats()], ["xlsx", "docx", "pdf"])

    def test_no_file_wanted(self):
        for phrase in NO_FILE:
            with self.subTest(phrase=phrase):
                self.assertEqual(detect(phrase), [])

    def test_a_new_format_needs_only_a_registration(self):
        original = _FORMATS.get("md")   # the real Markdown file format (builders/code.py) must survive this test
        register(ExportFormat(key="md", label="Markdown (.md)", extension="md", mime="text/markdown", builder=lambda d, l: b"# x",
                              terms=(r"\bmarkdown\b",), offer_by_default=False, order=60))
        try:
            self.assertEqual([f.key for f in detect("export it as markdown")], ["md"])
            self.assertIn("md", [f.key for f in all_formats()])
            self.assertEqual([f.key for f in detect("export it")], ["xlsx", "docx", "pdf"])
        finally:
            _FORMATS.pop("md", None)
            if original is not None:
                _FORMATS["md"] = original

    def test_route_name_matches_the_marker_used_to_hide_links(self):
        self.assertIn(EXPORT_PATH_MARK, reverse("ai-message-export", args=[1, "xlsx"]))


class RouterSeesOnlyTheQuestionTests(SimpleTestCase):
    """A file request must never change how the query engine routes the data question inside it."""

    BASES = ["what is my total balance", "list the paid salaries for year 2026", "show my certificates", "list my fixed assets",
             "gold price 21k", "exchange rate usd", "اعرض رواتب 2026", "liste mes salaires 2026", "Gehälter 2026 auflisten"]
    SUFFIXES = [", save it as csv", " and export it in excel sheet", " and download as word", ", export to pdf", " وصدرها اكسل", " et exporte en excel", " und als PDF herunterladen"]

    def test_routing_is_identical_with_and_without_the_file_request(self):
        for base in self.BASES:
            expected = route(base)
            for suffix in self.SUFFIXES:
                with self.subTest(base=base, suffix=suffix):
                    got = route(strip_export_terms(base + suffix))
                    self.assertEqual((got.status, got.capability), (expected.status, expected.capability))

    def test_bare_follow_ups_route_to_nothing(self):
        for phrase in ("export it to excel", "and as a pdf?", "give me this as word", "export it", "احفظ ده في ملف اكسل", "exporte ça en pdf", "als word herunterladen"):
            with self.subTest(phrase=phrase):
                self.assertEqual(route(strip_export_terms(phrase)).status, "none")


class DocumentTests(SimpleTestCase):
    def test_blocks(self):
        blocks = parse_markdown(SAMPLE)
        self.assertEqual([b.kind for b in blocks], ["heading", "table", "bullet", "text", "text"])
        table = blocks[1]
        self.assertEqual(table.header[:3], ["Month", "Company", "Expected"])
        self.assertEqual(len(table.rows), 3)
        self.assertEqual(blocks[2].text, "First note")
        self.assertEqual(blocks[3].text, "2. Second step".replace("2", "1"))

    def test_link_lines_never_reach_a_file(self):
        self.assertNotIn("export", strip_export_lines(SAMPLE))
        self.assertIn("[normal link](https://example.com)", strip_export_lines("[normal link](https://example.com)"))
        self.assertEqual(plain("[a](/x) **b** `c`"), "a b c")

    def test_ragged_and_headerless_tables(self):
        blocks = parse_markdown("| a | b | c |\n| 1 | 2 |")
        self.assertEqual(blocks[0].header, [])
        self.assertEqual([len(r) for r in blocks[0].rows], [3, 3])

    def test_numbers_move_their_unit_to_the_header(self):
        table = typed_table(parse_markdown(SAMPLE)[1])
        self.assertEqual(table.header[3], "Paid (EGP)")
        self.assertEqual(table.rows[0][3], 87643.86)
        self.assertEqual(table.decimals[3], 2)
        self.assertEqual(table.bold, [False, False, True])
        self.assertFalse(table.numeric[0])
        self.assertFalse(table.numeric[5])           # "+9.55%", "n/a", "" mix: stays text
        self.assertEqual(parse_number("+9.55%"), (9.55, "%", 2))
        self.assertIsNone(parse_number("Jan 2026"))


class BuilderTests(SimpleTestCase):
    def build(self, key, text=SAMPLE, lang="en"):
        return get_format(key).builder(doc_from(text, lang), lang)

    def test_xlsx_has_real_numbers_bold_total_and_no_links(self):
        ws = load_workbook(BytesIO(self.build("xlsx"))).active
        cells = {c.value: c for row in ws.iter_rows() for c in row if c.value is not None}
        self.assertEqual(cells["Paid (EGP)"].font.bold, True)
        self.assertEqual(cells[87643.86].number_format, "#,##0.00")
        self.assertEqual(cells[175644.36].font.bold, True)
        self.assertNotIn("Download as Excel (.xlsx)", cells)
        self.assertIn("• First note", cells)
        self.assertIn("Plain paragraph with bold word.", cells)

    def test_docx_has_table_text_and_rtl_for_arabic(self):
        doc = Document(BytesIO(self.build("docx")))
        self.assertEqual(doc.tables[0].rows[1].cells[3].text, "87,643.86 EGP")
        self.assertEqual(doc.tables[0].rows[-1].cells[0].text, "Grand Total")
        self.assertTrue(any(p.text == "First note" for p in doc.paragraphs))
        self.assertNotIn("w:bidi", Document(BytesIO(self.build("docx"))).element.xml)
        ar = Document(BytesIO(self.build("docx", AR_SAMPLE, "ar")))
        self.assertIn("w:bidi", ar.element.xml)
        self.assertEqual(ar.tables[0].rows[0].cells[0].text, "الشهر")

    def test_pdf_is_a_valid_pdf_for_english_and_arabic_and_wide_tables_go_landscape(self):
        for text, lang in ((SAMPLE, "en"), (AR_SAMPLE, "ar")):
            data = self.build("pdf", text, lang)
            self.assertTrue(data.startswith(b"%PDF-"))
            self.assertGreater(data.count(b"/Type /Page"), 0)
        wide = "| " + " | ".join("abcdefgh") + " |\n|" + "---|" * 8 + "\n| " + " | ".join("12345678") + " |"
        media = re.search(rb"/MediaBox \[ 0 0 ([\d.]+) ([\d.]+) \]", self.build("pdf", wide))
        self.assertGreater(float(media.group(1)), float(media.group(2)))

    def test_csv_numbers_and_utf8_bom(self):
        raw = self.build("csv").decode("utf-8")
        self.assertTrue(raw.startswith("﻿"))
        rows = list(csv.reader(StringIO(raw.lstrip("﻿"))))
        header = next(r for r in rows if "Paid (EGP)" in r)
        self.assertEqual(rows[rows.index(header) + 1][3], "87643.86")
        self.assertIn("الشهر", self.build("csv", AR_SAMPLE, "ar").decode("utf-8"))

    def test_txt_is_readable_and_link_free(self):
        text = self.build("txt").decode("utf-8")
        self.assertIn("Month | Company | Expected (EGP)".replace("Expected (EGP)", "Expected"), text)
        self.assertIn("- First note", text)
        self.assertNotIn("/export/", text)

    def test_empty_answer_still_builds_every_format(self):
        for fmt in all_formats():
            with self.subTest(fmt=fmt.key):
                self.assertTrue(fmt.builder(doc_from("", "en"), "en"))
