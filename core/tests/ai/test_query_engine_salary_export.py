"""Salary list + Excel export: the chat answer carries a working download link (0 LLM calls), the
endpoint builds an owner-scoped .xlsx, and the model path (direct answers off) gets the full-year list
and the same link instead of the capped 20-entry window."""

import subprocess
import shutil
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from django.conf import settings
from django.test import SimpleTestCase, TestCase
from openpyxl import load_workbook

from core.integrations.ai_provider import OllamaProvider
from core.models import Company, SalaryEntry
from core.services.ai.tools.salary_answers import MAX_YEAR_ENTRIES, SALARY_INSTRUCTIONS, add_salary_answers
from . import qe_fixtures as fx

NO_LLM = patch.object(OllamaProvider, "generate", side_effect=AssertionError("LLM must not be called"))
EXPORT_URL = "/api/salary/export/"
ASKS = [
    ("please list the paid salaries for year 2026, and export it in excel sheet", "en"),
    ("list my salaries 2026 and download them as an excel file", "en"),
    ("اعرض رواتب 2026 وصدرها في ملف اكسل", "ar"),
    ("liste mes salaires 2026 et exporte en excel", "fr"),
    ("Gehälter 2026 auflisten und als Excel herunterladen", "de"),
]


def _sheet(response):
    return load_workbook(BytesIO(response.content)).active


class SalaryExportChatTests(TestCase):
    def setUp(self):
        fx.build(self)

    def test_list_and_export_answers_instantly_with_link_and_table(self):
        for q, lang in ASKS:
            with self.subTest(q=q), NO_LLM as gen:
                data = fx.ask(self, q, conversation=False)
                text = data["message"]["content"]
                self.assertIn("87,643.86", text)
                self.assertIn(f"]({EXPORT_URL}?years=2026&lang={lang})", text)
                self.assertEqual(data["message"]["tool_calls"][0]["llm_calls"], 0)
                gen.assert_not_called()

    def test_no_export_words_no_link(self):
        with NO_LLM:
            text = fx.ask(self, "list the paid salaries for year 2026")["message"]["content"]
        self.assertIn("87,643.86", text)
        self.assertNotIn(EXPORT_URL, text)

    def test_nothing_to_export_no_link(self):
        with NO_LLM:
            text = fx.ask(self, "list salaries 2019 and export to excel")["message"]["content"]
        self.assertNotIn(EXPORT_URL, text)

    def test_multi_year_link_names_every_year(self):
        co = Company.objects.get(owner=self.user)
        SalaryEntry.objects.create(company=co, year=2025, month="May", paid="1000", expected="1000")
        with NO_LLM:
            text = fx.ask(self, "salaries for 2025 and 2026 export to excel")["message"]["content"]
        self.assertIn(f"{EXPORT_URL}?years=2025,2026&lang=en", text)


class SalaryExportEndpointTests(TestCase):
    def setUp(self):
        fx.build(self)
        co = Company.objects.get(owner=self.user)
        SalaryEntry.objects.create(company=co, year=2025, month="December", paid="80000", expected="80000", bonus="500")

    def test_requires_login(self):
        self.client.logout()
        self.assertIn(self.client.get(EXPORT_URL).status_code, (401, 403))

    def test_year_file_has_rows_and_total_and_no_other_tenant_data(self):
        r = self.client.get(EXPORT_URL, {"years": "2026", "lang": "en"})
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])
        self.assertIn("Salaries_2026.xlsx", r["Content-Disposition"])
        rows = [[c.value for c in row] for row in _sheet(r).iter_rows()]
        self.assertEqual(rows[0][:3], ["Year", "Month", "Company"])
        self.assertEqual(rows[1][:3], [2026, "January", "Acme"])
        self.assertEqual(rows[1][3:5], [90000.0, 87643.86])
        self.assertEqual(rows[-1][2:5], ["Grand Total", 90000.0, 87643.86])
        self.assertEqual(len(rows), 3)
        self.assertNotIn(111.0, [v for row in rows for v in row])

    def test_no_years_exports_everything_oldest_first(self):
        r = self.client.get(EXPORT_URL)
        self.assertIn("Salaries_All.xlsx", r["Content-Disposition"])
        rows = [[c.value for c in row] for row in _sheet(r).iter_rows()]
        self.assertEqual([row[0] for row in rows[1:3]], [2025, 2026])
        self.assertEqual(rows[-1][4], 167643.86)
        self.assertEqual(rows[-1][5], 500.0)

    def test_arabic_headers_and_rtl(self):
        ws = _sheet(self.client.get(EXPORT_URL, {"years": "2026", "lang": "ar"}))
        self.assertEqual(ws["A1"].value, "السنة")
        self.assertTrue(ws.sheet_view.rightToLeft)

    def test_unknown_language_falls_back_to_english_and_junk_years_are_ignored(self):
        ws = _sheet(self.client.get(EXPORT_URL, {"years": "abc,20x6,2026", "lang": "xx"}))
        self.assertEqual(ws["A1"].value, "Year")
        self.assertEqual(ws.max_row, 3)

    def test_nothing_to_export_is_404(self):
        self.assertEqual(self.client.get(EXPORT_URL, {"years": "2019"}).status_code, 404)

    def test_other_tenant_gets_only_own_rows(self):
        self.client.force_login(self.other)
        rows = [[c.value for c in row] for row in _sheet(self.client.get(EXPORT_URL, {"years": "2026"})).iter_rows()]
        self.assertEqual(rows[1][4], 111.0)


class SalaryModelPathTests(TestCase):
    """Direct answers OFF: the tool payload must still carry the whole year and the link."""

    def setUp(self):
        fx.build(self)
        co = Company.objects.get(owner=self.user)
        for i, month in enumerate(("February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December")):
            SalaryEntry.objects.create(company=co, year=2026, month=month, paid=str(1000 + i), expected="1000")
        co2 = Company.objects.create(owner=self.user, name="Beta", display_name="Beta")
        for month in ("January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"):
            SalaryEntry.objects.create(company=co2, year=2026, month=month, paid="5", expected="5")

    def test_year_list_is_not_capped_at_20_and_link_is_present(self):
        res = {"salary": {"latest_salary_entry": None}}
        add_salary_answers(res, self.user, "list the paid salaries for year 2026 and export it in excel sheet")
        salary = res["salary"]
        self.assertEqual(len(salary["requested_year_entries"]), 24)
        self.assertGreater(len(salary["requested_year_entries"]), 20)
        self.assertEqual(salary["export_link_markdown"], f"[Download as Excel (.xlsx)]({EXPORT_URL}?years=2026&lang=en)")

    def test_year_named_without_list_or_export_words_still_gets_the_full_year(self):
        res = {"salary": {}}
        add_salary_answers(res, self.user, "salary 2026")
        self.assertEqual(len(res["salary"]["requested_year_entries"]), 24)

    def test_no_year_means_no_list_and_a_link_for_everything(self):
        res = {"salary": {}}
        add_salary_answers(res, self.user, "what is my average salary")
        self.assertNotIn("requested_year_entries", res["salary"])
        self.assertEqual(res["salary"]["export_link_markdown"], f"[Download as Excel (.xlsx)]({EXPORT_URL}?lang=en)")

    def test_arabic_query_gets_arabic_link_text(self):
        res = {"salary": {}}
        add_salary_answers(res, self.user, "رواتب 2026")
        self.assertIn("lang=ar", res["salary"]["export_link_markdown"])

    def test_year_list_only_for_the_owner(self):
        res = {"salary": {}}
        add_salary_answers(res, self.other, "list salaries 2026")
        entries = res["salary"]["requested_year_entries"]
        self.assertEqual(len(entries), 1)
        self.assertIn("111.00", entries[0]["paid_formatted"])

    def test_cap_and_instruction_text(self):
        self.assertGreaterEqual(MAX_YEAR_ENTRIES, 100)
        self.assertIn("export_link_markdown", SALARY_INSTRUCTIONS)
        self.assertIn("requested_year_entries", SALARY_INSTRUCTIONS)


class ChatMarkdownLinkTests(SimpleTestCase):
    """static/js/ai/ai_markdown.js renders app download links and nothing external."""

    def _render(self, text):
        node = shutil.which("node")
        if not node:
            self.skipTest("node not available")
        js = Path(settings.BASE_DIR, "static/js/ai/ai_markdown.js").read_text(encoding="utf-8")
        script = js + "\nprocess.stdout.write(_renderMarkdown(process.argv[1]));"
        return subprocess.run([node, "-e", script, text], capture_output=True, text=True, check=True).stdout

    def test_api_download_link_becomes_anchor(self):
        html = self._render(f"[Download as Excel (.xlsx)]({EXPORT_URL}?years=2026&lang=en)")
        self.assertIn('<a class="ai-ws-link" href="/api/salary/export/?years=2026&amp;lang=en" download', html)
        self.assertIn("Download as Excel (.xlsx)</a>", html)

    def test_external_and_script_links_stay_text(self):
        html = self._render("[a](https://evil.example/x) [b](javascript:alert(1)) [c](//evil.example)")
        self.assertNotIn("<a ", html)


class MultiYearTimeParsingTests(SimpleTestCase):
    """'2025 and 2026' used to keep only the first year, which silently dropped the second from lists and exports."""

    def _years(self, text):
        from datetime import date
        from core.services.ai.query_engine.lexicon import norm
        from core.services.ai.query_engine.timespec import parse_time

        months = parse_time(norm(text), date(2026, 10, 12)).months
        return sorted({y for y, _ in months}), len(months)

    def test_year_lists_and_spans(self):
        for text, years in (("salaries for 2025 and 2026", [2025, 2026]), ("salaries 2024, 2025 and 2026", [2024, 2025, 2026]),
                            ("from 2025 to 2026", [2025, 2026]), ("رواتب 2025 و 2026", [2025, 2026]),
                            ("salary 2026", [2026]), ("in 2025", [2025])):
            with self.subTest(text=text):
                got, n = self._years(text)
                self.assertEqual(got, years)
                self.assertEqual(n, 12 * len(years))

    def test_month_ranges_are_untouched(self):
        self.assertEqual(self._years("jun 2026 to sep 2026"), ([2026], 4))
        self.assertEqual(self._years("2026-01 to 2026-03"), ([2026], 3))
