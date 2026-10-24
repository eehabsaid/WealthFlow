"""Guide files (PDF / Word) and refusals: RTL line order, outline + contents + page numbers, Word heading levels + TOC field,
no repeated purpose paragraph, no file for a refusal, the html matcher, and the i18n scope refusal in four languages."""

from __future__ import annotations

import json
import re
import zipfile
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.test import SimpleTestCase, TestCase

from core.reports.pdf_rtl import rtl_markup, visual, wrap_lines
from core.services.ai import scope_text
from core.services.ai.domain_scope import is_no_info, no_info
from core.services.ai.domains import APP, BUSINESS
from core.services.ai.exports.builders import docx as docx_builder, pdf as pdf_builder
from core.services.ai.exports.builders.common import guide_blocks
from core.services.ai.exports.document import ExportDocument, parse_markdown
from core.services.ai.exports.intent import formats_in

PURPOSE = "الغرض: هذه الصفحة تتيح لك إدارة الحسابات المصرفية وعرض الأرصدة وتسجيل الإيداعات والسحوبات ومراجعة الحركات السابقة (Bank Accounts) بشكل واضح."
PURPOSE_EN = "Purpose: this page lets you manage bank accounts and review balances and past transactions in a clear way for every user."


def guide_markdown(text: str, chapters: int = 6) -> str:
    md = ""
    for i in range(1, chapters + 1):
        md += f"## Chapter {i}\n\n{text}\n\n### Tab {i}\n\n1. Open the page {i}.\n2. Press save.\n\n#### Dialog {i}\n\n- {text}\n\n"
    return md


def make_doc(md: str, lang: str) -> ExportDocument:
    return ExportDocument(title="Guide", generated="2026", lang=lang, blocks=parse_markdown(md), source=md)


class RtlWrapTest(SimpleTestCase):
    FACES = ("Helvetica", "Helvetica-Bold")

    def test_lines_follow_the_logical_order_and_each_fits(self):
        lines = wrap_lines([(PURPOSE * 3, False)], self.FACES, 10, 200)
        self.assertGreater(len(lines), 2)
        flat = " ".join(w for line in lines for w, _ in line)
        self.assertEqual(flat, " ".join((PURPOSE * 3).split()))   # the words stay in logical order across lines
        self.assertTrue(lines[0][0][0].startswith("الغرض"))

    def test_a_latin_term_keeps_its_letters_in_order_and_the_markup_has_one_br_per_break(self):
        markup = rtl_markup([("افتح (Bank Accounts) الآن", False)], self.FACES, 10, 500)
        self.assertEqual(markup.count("<br/>"), 0)
        self.assertIn("Bank Accounts", markup)
        self.assertIn("Bank", visual("(Bank Accounts)"))

    def test_bold_words_stay_bold(self):
        self.assertIn("<b>", rtl_markup([("مرحبا ", False), ("بالعالم", True)], self.FACES, 500, 500))


class GuidePdfTest(SimpleTestCase):
    def pdf(self, lang: str):
        import pypdf

        text = PURPOSE if lang == "ar" else PURPOSE_EN
        return pypdf.PdfReader(BytesIO(pdf_builder._build(make_doc(guide_markdown(text, 12), lang), lang)))

    def test_guide_pdf_has_outline_contents_and_page_numbers_in_both_directions(self):
        for lang in ("en", "ar"):
            reader = self.pdf(lang)
            titles = [o.title for o in reader.outline if not isinstance(o, list)]
            self.assertEqual(len(titles), 12, lang)                      # chapters
            self.assertTrue(any(isinstance(o, list) and len(o) == 1 for o in reader.outline), lang)   # their tabs
            first = reader.pages[0].extract_text()
            self.assertTrue(re.search(r"\d+\s*$", first.strip()), lang)   # footer page number
            self.assertGreaterEqual(len(reader.pages), 2)

    def test_non_guide_pdf_is_unchanged_in_shape(self):
        reader = __import__("pypdf").PdfReader(BytesIO(pdf_builder._build(make_doc("## A\n\ntext\n", "en"), "en")))
        self.assertEqual(len(reader.outline), 0)


class GuideDocxTest(SimpleTestCase):
    def build(self, lang: str):
        import docx

        data = docx_builder._build(make_doc(guide_markdown(PURPOSE if lang == "ar" else PURPOSE_EN), lang), lang)
        return docx.Document(BytesIO(data)), zipfile.ZipFile(BytesIO(data))

    def test_levels_toc_field_and_update_fields(self):
        for lang in ("en", "ar"):
            document, z = self.build(lang)
            counts = {}
            for p in document.paragraphs:
                counts[p.style.name] = counts.get(p.style.name, 0) + 1
            self.assertEqual((counts["Heading 1"], counts["Heading 2"], counts["Heading 3"]), (6, 6, 6), lang)
            xml = z.read("word/document.xml").decode()
            self.assertIn('TOC \\o "1-2"', xml)
            self.assertIn("updateFields", z.read("word/settings.xml").decode())
            self.assertEqual("w:bidi" in xml, lang == "ar")

    def test_updatefields_sits_before_compat(self):
        _, z = self.build("en")
        settings_xml = z.read("word/settings.xml").decode()
        self.assertLess(settings_xml.index("updateFields"), settings_xml.index("<w:compat"))

    def test_a_non_guide_keeps_its_heading_levels(self):
        import docx

        data = docx_builder._build(make_doc("# T\n\n## A\n\ntext\n\n### B\n\ntext\n", "en"), "en")
        names = [p.style.name for p in docx.Document(BytesIO(data)).paragraphs]
        self.assertIn("Heading 2", names)


class RepeatedPurposeTest(SimpleTestCase):
    def test_a_paragraph_repeated_under_another_tab_is_written_once_but_steps_stay(self):
        md = "## One\n\nSame purpose text here.\n\n1. Step\n\n### Tab\n\nSame purpose text here.\n\n1. Step\n"
        blocks = guide_blocks(make_doc(md, "en"))
        self.assertEqual(sum(b.text == "Same purpose text here." for b in blocks), 1)
        self.assertEqual(sum(b.text.startswith("1.") for b in blocks), 2)


class RefusalAndHtmlTest(TestCase):
    def test_the_word_html_alone_does_not_make_a_file_request_but_a_page_does(self):
        self.assertEqual([f.key for f in formats_in("what is html used for in the app?")], [])
        self.assertIn("html", [f.key for f in formats_in("write me an html login page")])
        self.assertIn("html", [f.key for f in formats_in("save the guide as guide.html")])
        self.assertIn("html", [f.key for f in formats_in("اعمل لي دليل بصيغة html")])

    def test_a_refusal_is_recognised_in_every_language_and_never_gets_a_file(self):
        from core.views.ai_chat.ai_chat_core_views.response_finalizer import is_non_answer

        for lang in scope_text.LANGS:
            for domain in (BUSINESS, APP, ""):
                self.assertTrue(is_non_answer(scope_text.refusal(lang, domain)), (lang, domain))
        self.assertFalse(is_non_answer("| a | b |\n|---|---|\n| 1 | 2 |"))
        self.assertFalse(is_non_answer("Your total is 100 across 3 accounts. " * 30))


class ScopeRefusalTextTest(TestCase):
    def test_it_names_the_domains_from_the_ui_and_says_where_to_switch(self):
        for lang in scope_text.LANGS:
            ui = json.loads((Path(settings.BASE_DIR) / "static" / "i18n" / f"{lang}.json").read_text(encoding="utf-8"))
            for domain, other in ((BUSINESS, APP), (APP, BUSINESS)):
                msg = scope_text.refusal(lang, domain)
                key = {BUSINESS: "ai_domain_business_data", APP: "ai_domain_app_features"}
                self.assertIn(ui[key[domain]], msg)
                self.assertIn(ui[key[other]], msg)
                self.assertNotIn("{", msg)

    def test_no_info_follows_the_question_language(self):
        self.assertIn("المحادثة", no_info("ما هو رصيدي", BUSINESS))
        self.assertTrue(is_no_info(no_info("what is my balance", APP)))
