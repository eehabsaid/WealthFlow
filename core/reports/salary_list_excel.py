"""Salary list -> .xlsx for the AI chat download link and GET /api/salary/export/.

One sheet, one row per salary entry (year, month, company, expected, paid, bonus) plus a total row.
Owner-scoped, amounts in the user's base currency label, headers in en / ar / fr / de (same label
tables the query engine uses, so the file matches the chat table). Totals are written as values so
every viewer (including phone previews that do not calculate formulas) shows them.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from core.services.ai.query_engine.labels import month_label, t

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_MONEY = "#,##0.00"
_SUPPORTED_LANGS = ("en", "ar", "fr", "de")


def parse_years(raw: str) -> list[int]:
    """'2026,2025' -> [2025, 2026]; ignores anything that is not a 4-digit year."""
    years = {int(p) for p in str(raw or "").replace(";", ",").split(",") if p.strip().isdigit() and len(p.strip()) == 4}
    return sorted(years)


def salary_rows(user: Any, years: Iterable[int] = ()) -> list[tuple[int, int, Any]]:
    """[(year, month_number, SalaryEntry)] oldest first; month_number 0 when the stored month name is unknown."""
    from core.models import SalaryEntry
    from core.services.ai.providers.salary_provider.constants import MONTH_NAME_TO_INT

    qs = SalaryEntry.objects.filter(company__owner=user).select_related("company")
    wanted = set(years)
    if wanted:
        qs = qs.filter(year__in=wanted)
    rows = [(e.year, MONTH_NAME_TO_INT.get(str(e.month or "").strip().lower(), 0), e) for e in qs]
    return sorted(rows, key=lambda r: (r[0], r[1], r[2].id))


def build_salary_workbook(user: Any, years: Iterable[int], lang: str) -> tuple[BytesIO, int]:
    """(xlsx buffer, number of entry rows). Raises nothing for an empty selection: the caller checks the count."""
    from core.services.shared.base_currency import get_user_base_code

    lang = lang if lang in _SUPPORTED_LANGS else "en"
    cur = get_user_base_code(user)
    rows = salary_rows(user, years)
    wb = Workbook()
    ws = wb.active
    ws.title = t(lang, "sal_sheet")[:31]
    ws.sheet_view.rightToLeft = lang == "ar"
    headers = [t(lang, "col_year"), t(lang, "col_month"), t(lang, "col_company"),
               f"{t(lang, 'col_expected')} ({cur})", f"{t(lang, 'col_paid')} ({cur})", f"{t(lang, 'col_bonus')} ({cur})"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sums = [0.0, 0.0, 0.0]
    for year, month, e in rows:
        vals = [float(e.expected or 0), float(e.paid or 0), float(e.bonus or 0)]
        sums = [a + b for a, b in zip(sums, vals)]
        label = month_label(lang, year, month)[: -len(str(year)) - 1] if month else str(e.month or "")
        ws.append([year, label, e.company.name if e.company else "", *vals])
    ws.append(["", "", t(lang, "grand"), *sums])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=6):
        for cell in row:
            cell.number_format = _MONEY
    for col, width in zip("ABCDEF", (8, 16, 28, 20, 20, 20)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf, len(rows)
