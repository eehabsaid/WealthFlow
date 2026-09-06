"""Per-year row group builder for the salary sheet builder.

Pure code motion from the original excel_salary_builder.py (the body of the
per-year loop inside build_salary_sheet) — same logic, same order of
operations, only relocated for file-size compliance.
"""

from openpyxl.styles import Font, Border, Side

from core.reports.excel_formatting_helpers import FMT_EGP, _thin, _fill, _align
from .constants import YEAR_ROW_HT


def build_year_group(ws, row, year, year_entries, name, cols, last_col, has_bonus, is_first_group):
    """
    Write one year's label row, data rows, and per-year total row.
    `is_first_group` mirrors the original `len(total_rows) == 0` check.
    Returns (row_after_group, total_row_index).
    """
    fill_dark_blue = _fill("FF1F497D")
    fill_red_data = _fill("FFC0504D")

    yr_row = row
    ws.row_dimensions[yr_row].height = YEAR_ROW_HT
    yr_merge = f"A{yr_row}:{last_col}{yr_row}"
    try:
        ws.merge_cells(yr_merge)
    except Exception:
        pass
    yc = ws.cell(row=yr_row, column=1, value=int(year))
    yc.font = Font(bold=True, size=18, name="Times New Roman")
    yc.fill = fill_dark_blue
    yc.alignment = _align("center")
    if has_bonus:
        yc.border = Border()
    else:
        yc.border = Border(top=Side(style="thin"), bottom=Side(style="thin"))
    row += 1

    data_start = row
    for entry in year_entries:
        for c in range(1, cols + 1):
            ws.cell(row=row, column=c).fill = fill_red_data
            ws.cell(row=row, column=c).font = Font(size=11, name="Arial")
            ws.cell(row=row, column=c).border = _thin()

        ws.cell(row=row, column=1, value=int(entry.year))
        ws.cell(row=row, column=2, value=str(entry.month))
        ws.cell(row=row, column=3, value=float(entry.expected))
        ws.cell(row=row, column=3).number_format = FMT_EGP
        ws.cell(row=row, column=4, value=float(entry.paid))
        ws.cell(row=row, column=4).number_format = FMT_EGP

        if name in ("NTG", "Giza Systems", "Giza Systems (2)"):
            rem = f"=D{row}-C{row}"
        else:
            rem = f"=IF(C{row}>D{row},C{row}-D{row},0)"
        ws.cell(row=row, column=5, value=rem)
        ws.cell(row=row, column=5).number_format = FMT_EGP

        if has_bonus:
            bonus_val = float(getattr(entry, "bonus", 0) or 0)
            ws.cell(row=row, column=6, value=bonus_val)
            ws.cell(row=row, column=6).number_format = FMT_EGP
        row += 1

    data_end = row - 1

    paid_count = sum(1 for e in year_entries if float(e.paid) > 0)

    for c in range(1, cols + 1):
        tc = ws.cell(row=row, column=c)
        tc.fill = fill_dark_blue
        tc.font = Font(bold=True, size=11, name="Arial")
        tc.alignment = _align("center")
        tc.border = _thin()

    ws.cell(row=row, column=1, value="Total")
    if name == "NTG" and is_first_group:
        ws.cell(row=row, column=2, value=paid_count)
    else:
        ws.cell(
            row=row,
            column=2,
            value=f'=COUNTIF(D{data_start}:D{data_end}, "<> 0.00")',
        )
    ws.cell(row=row, column=3, value=f"=SUM(C{data_start}:C{data_end})")
    ws.cell(row=row, column=3).number_format = FMT_EGP
    ws.cell(row=row, column=4, value=f"=SUM(D{data_start}:D{data_end})")
    ws.cell(row=row, column=4).number_format = FMT_EGP

    if name == "NTG":
        ws.cell(row=row, column=5, value=f"=D{row}-C{row}")
    else:
        ws.cell(row=row, column=5, value=f"=SUM(E{data_start}:E{data_end})")
    ws.cell(row=row, column=5).number_format = FMT_EGP

    if has_bonus:
        ws.cell(row=row, column=6, value=f"=SUM(F{data_start}:F{data_end})")
        ws.cell(row=row, column=6).number_format = FMT_EGP

    total_row = row
    row += 1

    return row, total_row
