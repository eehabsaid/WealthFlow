"""Sheet frame, header row, and title-row setup for the salary sheet builder.

Pure code motion from the original excel_salary_builder.py (the top part of
build_salary_sheet, before the per-year loop) — same logic, same order of
operations, only relocated for file-size compliance.
"""

from openpyxl.styles import Font, Border, Side

from core.reports.excel_formatting_helpers import _thin, _fill, _align
from .constants import SALARY_COL_WIDTHS, SALARY_FREEZE

DEFAULT_WIDTHS = {"A": 13.7, "B": 16.1, "C": 16.0, "D": 19.3, "E": 15.6, "F": 14.3}
GREY = "FF7F7F7F"
RED_TTL = "FFFF0000"


def setup_sheet_frame(ws, name):
    """Row heights, column widths, and freeze panes for rows 1-3."""
    ws.row_dimensions[1].height = 14.25
    ws.row_dimensions[2].height = 14.25
    ws.row_dimensions[3].height = 20.25

    widths = SALARY_COL_WIDTHS.get(name, DEFAULT_WIDTHS)
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    fp = SALARY_FREEZE.get(name)
    if fp:
        ws.freeze_panes = fp


def build_header_row(ws, has_bonus):
    """Row 1: column headers."""
    fill_dark_blue = _fill("FF1F497D")
    hdrs = [
        "Year",
        "Month",
        "Expected",
        "Paid (Salary + Bonus)" if has_bonus else "Paid",
        "Remaining",
    ]
    if has_bonus:
        hdrs.append("Bonus")

    for c, h in enumerate(hdrs, 1):
        cell = ws.cell(row=1, column=c, value=h)
        if has_bonus:
            cell.font = Font(bold=True, italic=True, size=11, name="Arial")
        else:
            cell.font = Font(bold=False, italic=True, size=11, name="Arial", color=GREY)
        cell.fill = fill_dark_blue
        cell.alignment = _align("center")
        cell.border = _thin()


def build_title_row(ws, last_col, has_bonus):
    """Rows 2-3 merged: the 'Salary Details' title banner."""
    fill_black = _fill("FF000000")
    ws.merge_cells(f"A2:{last_col}3")
    c2 = ws.cell(row=2, column=1, value=" Salary Details")
    c2.font = Font(bold=True, size=18, name="Times New Roman", color=RED_TTL)
    c2.fill = fill_black
    c2.alignment = _align("center")
    if has_bonus:
        c2.border = Border(top=Side(style="thin"))
    else:
        c2.border = Border(bottom=Side(style="thin"))
