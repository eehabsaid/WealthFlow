"""Grand-total row builder for the salary sheet builder.

Pure code motion from the original excel_salary_builder.py (the tail of
build_salary_sheet, after the per-year loop) — same logic, same order of
operations, only relocated for file-size compliance.
"""

from openpyxl.styles import Font

from core.reports.excel_formatting_helpers import FMT_EGP, _thin, _fill

LABEL_MAP = {
    "NTG": "Total",
    "Giza Systems": "Total",
    "Giza Systems (2)": "Total",
    "Giza Systems (3)": "Total",
    "ElSeweedy Technology": "Total",
    "Dedalus": "Total",
    "Globemed": "Total",
}


def build_grand_total(ws, row, total_rows, name, cols, has_bonus):
    """Write the final grand-total row summing every per-year total row."""
    fill_black = _fill("FF000000")
    sr = row
    label = LABEL_MAP.get(name, "Total")

    for c in range(1, cols + 1):
        sc = ws.cell(row=sr, column=c)
        sc.fill = fill_black
        sc.font = Font(bold=True, size=11, name="Arial", color="FFFFFFFF")
        sc.border = _thin()

    ws.cell(row=sr, column=1, value=label)
    b_ref = "+".join(f"B{r}" for r in total_rows)
    ws.cell(row=sr, column=2, value=f"={b_ref}")
    ws.cell(row=sr, column=2).number_format = "0" if name != "NTG" else "General"
    c_ref = "+".join(f"C{r}" for r in total_rows)
    d_ref = "+".join(f"D{r}" for r in total_rows)
    ws.cell(row=sr, column=3, value=f"={c_ref}")
    ws.cell(row=sr, column=3).number_format = FMT_EGP
    ws.cell(row=sr, column=4, value=f"={d_ref}")
    ws.cell(row=sr, column=4).number_format = FMT_EGP
    ws.cell(row=sr, column=5, value=f"=D{sr}-C{sr}")
    ws.cell(row=sr, column=5).number_format = FMT_EGP

    if has_bonus:
        f_ref = "+".join(f"F{r}" for r in total_rows)
        ws.cell(row=sr, column=6, value=f"={f_ref}")
        ws.cell(row=sr, column=6).number_format = FMT_EGP

    ws.row_dimensions[sr].height = 21.0 if name == "NTG" else None

    return sr
