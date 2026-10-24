"""Word (.docx): title block, then headings / paragraphs / bullets / tables in answer order.
Bold emphasis is kept; Arabic answers get right-to-left paragraphs and mirrored tables."""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ..document import ExportDocument, is_bold_cell, plain, runs
from ..registry import ExportFormat, register
from .common import guide_blocks, intro_lines, is_wide
from .docx_toc import add_contents, request_field_update, word_heading_level
from .guide_html import is_guide

_HEADER_SHADE = "D9E2F3"
_TABLE_STYLE = "Table Grid"
_MAX_HEADING_LEVEL = 4


def _rtl_paragraph(paragraph) -> None:
    # w:bidi alone starts the text at the right edge; an explicit jc="right" is mirrored by Word / LibreOffice for a bidi paragraph
    paragraph._p.get_or_add_pPr().append(OxmlElement("w:bidi"))


def _add_runs(paragraph, text: str, rtl: bool, force_bold: bool = False) -> None:
    for chunk, bold in runs(text) or [(plain(text), False)]:
        run = paragraph.add_run(chunk)
        run.bold = bold or force_bold or None
        if rtl:
            run._r.get_or_add_rPr().append(OxmlElement("w:rtl"))


def _shade(cell, fill: str) -> None:
    shade = OxmlElement("w:shd")
    shade.set(qn("w:val"), "clear")
    shade.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shade)


def _add_table(document, block, rtl: bool) -> None:
    width = len(block.header) if block.header else max((len(r) for r in block.rows), default=0)
    if not width:
        return
    table = document.add_table(rows=0, cols=width)
    table.style = _TABLE_STYLE
    if rtl:
        table._tbl.tblPr.append(OxmlElement("w:bidiVisual"))
    if block.header:
        cells = table.add_row().cells
        for cell, text in zip(cells, block.header):
            _add_runs(cell.paragraphs[0], text, rtl, force_bold=True)
            _shade(cell, _HEADER_SHADE)
    for row in block.rows:
        bold = any(is_bold_cell(c) for c in row)
        for cell, text in zip(table.add_row().cells, row):
            _add_runs(cell.paragraphs[0], text, rtl, force_bold=bold)
            if rtl:
                _rtl_paragraph(cell.paragraphs[0])
    document.add_paragraph()


def _build(doc: ExportDocument, lang: str) -> bytes:
    rtl = lang == "ar"
    document = Document()
    guide = is_guide(doc)
    if is_wide(doc):
        section = document.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    for role, text in intro_lines(doc):
        paragraph = document.add_heading(text, level=0) if role == "title" else document.add_paragraph()
        if role != "title":
            _add_runs(paragraph, text, rtl, force_bold=role == "question")
        if rtl:
            _rtl_paragraph(paragraph)
    blocks, contents_done = (guide_blocks(doc) if guide else doc.blocks), not guide
    for block in blocks:
        if not contents_done and block.kind == "heading" and block.level == 2:
            add_contents(document, blocks, lang, rtl)
            contents_done = True
        if block.kind == "table":
            _add_table(document, block, rtl)
            continue
        if block.kind == "heading":
            level = word_heading_level(block.level) if guide else min(max(block.level, 1), _MAX_HEADING_LEVEL)
            paragraph = document.add_heading(plain(block.text), level=level)
        else:
            paragraph = document.add_paragraph(style="List Bullet" if block.kind == "bullet" else None)
            _add_runs(paragraph, block.text, rtl)
        if rtl:
            _rtl_paragraph(paragraph)
    if guide:
        request_field_update(document)
    buf = BytesIO()
    document.save(buf)
    return buf.getvalue()


register(ExportFormat(
    key="docx", label="Word (.docx)", extension="docx",
    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", builder=_build, order=20,
    terms=(r"\bdocx?\b|\bms word\b|\bmicrosoft word\b|\bword (?:file|document|doc|format)\b|\b(?:to|in|as|into|via|with) (?:a |an )?word\b|\bword-?(?:datei|dokument)\b|\b(?:als|in|nach) word\b|\bfichier word\b|\bau format word\b|\ben word\b",
           r"(?:ملف|الي|الى|في|كملف|بصيغه|بصيغة|ب)\s*(?:ال)?(?:وورد|ورد)\b|\bوورد\b|\bدوكس\b"),
    weak_terms=(r"\bword\b", r"\bوورد\b|\bورد\b"),
))
