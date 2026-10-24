"""A user guide in Word: chapters are Heading 1, tabs Heading 2, dialogs Heading 3, with a real table-of-contents field over
Heading 1-2 (Word fills in the page numbers when the file is opened: updateFields) and a cached list so the contents is not
empty before that. The guide's markdown levels are title #, chapter ##, tab ###, dialog ####; the HTML guide still reads them."""

from __future__ import annotations

from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ..document import plain

TOC_INSTRUCTION = r'TOC \o "1-2" \h \z \u'


def word_heading_level(markdown_level: int) -> int:
    """## chapter -> 1, ### tab -> 2, #### and deeper (dialogs) -> 3."""
    return min(max(markdown_level - 1, 1), 3)


def _field_run(paragraph, kind: str = "", instruction: str = ""):
    run = paragraph.add_run()
    if instruction:
        node = OxmlElement("w:instrText")
        node.set(qn("xml:space"), "preserve")
        node.text = instruction
    else:
        node = OxmlElement("w:fldChar")
        node.set(qn("w:fldCharType"), kind)
    run._r.append(node)


def add_contents(document, blocks: list, lang: str, rtl: bool) -> None:
    """Contents title, the TOC field (cached entries inside it) and a page break."""
    from .guide_html_assets import LABELS

    title = document.add_paragraph()
    title.add_run(LABELS.get(lang, LABELS["en"])["contents"]).bold = True
    entries = [(b.level, plain(b.text)) for b in blocks if b.kind == "heading" and b.level in (2, 3)]
    paragraphs = [document.add_paragraph() for _ in entries or [None]]
    first, last = paragraphs[0], paragraphs[-1]
    _field_run(first, "begin")
    _field_run(first, instruction=TOC_INSTRUCTION)
    _field_run(first, "separate")
    for paragraph, (level, text) in zip(paragraphs, entries):
        paragraph.add_run(text)
        paragraph.paragraph_format.left_indent = None if level == 2 else 360 * 635
    _field_run(last, "end")
    for paragraph in [title, *paragraphs]:
        if rtl:
            paragraph._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def request_field_update(document) -> None:
    """Word refreshes the contents (and its page numbers) when the file is opened."""
    settings = document.settings.element
    node = OxmlElement("w:updateFields")
    node.set(qn("w:val"), "true")
    later = settings.find(qn("w:hdrShapeDefaults"))   # schema order: updateFields comes before these (Word rejects a wrong order)
    for tag in ("w:footnotePr", "w:endnotePr", "w:compat"):
        later = settings.find(qn(tag)) if later is None else later
    if later is not None:
        later.addprevious(node)
    else:
        settings.append(node)
